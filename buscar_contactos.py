"""
Script para buscar contactos y webs de puntos de venta en Argentina.
Usa la API de Google Places para buscar automáticamente cada tienda.

Requisitos:
    pip install openpyxl requests

Uso:
    python3 buscar_contactos.py --api-key TU_API_KEY --input archivo.xlsx --output resultado.xlsx

Obtener API key de Google Places:
    1. Ir a https://console.cloud.google.com/
    2. Crear proyecto → Habilitar "Places API"
    3. Crear credencial → API Key
    (Tiene capa gratuita de $200/mes ≈ 5000 búsquedas gratis)
"""

import argparse
import time
import json
import sys
import requests
import openpyxl
from pathlib import Path


PLACES_FIND_URL = "https://maps.googleapis.com/maps/api/place/findplacefromtext/json"
PLACES_DETAIL_URL = "https://maps.googleapis.com/maps/api/place/details/json"


def find_place(razon_social: str, domicilio: str, localidad: str, api_key: str) -> dict:
    """Busca un negocio en Google Places y retorna phone + website."""
    query = f"{razon_social}, {domicilio}, {localidad}, Argentina"

    # Step 1: Find place ID
    resp = requests.get(PLACES_FIND_URL, params={
        "input": query,
        "inputtype": "textquery",
        "fields": "place_id,name,formatted_address",
        "locationbias": "country:AR",
        "key": api_key,
    }, timeout=10)
    resp.raise_for_status()
    data = resp.json()

    candidates = data.get("candidates", [])
    if not candidates:
        # Retry with simpler query (just address + localidad)
        resp = requests.get(PLACES_FIND_URL, params={
            "input": f"{domicilio}, {localidad}, Argentina",
            "inputtype": "textquery",
            "fields": "place_id,name,formatted_address",
            "locationbias": "country:AR",
            "key": api_key,
        }, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        candidates = data.get("candidates", [])

    if not candidates:
        return {"phone": None, "website": None, "found": False}

    place_id = candidates[0]["place_id"]

    # Step 2: Get place details
    resp = requests.get(PLACES_DETAIL_URL, params={
        "place_id": place_id,
        "fields": "name,formatted_phone_number,international_phone_number,website",
        "key": api_key,
    }, timeout=10)
    resp.raise_for_status()
    detail = resp.json().get("result", {})

    phone = detail.get("international_phone_number") or detail.get("formatted_phone_number")
    website = detail.get("website")

    return {"phone": phone, "website": website, "found": True}


def process_excel(input_path: str, output_path: str, api_key: str, delay: float = 0.3, start_row: int = 0):
    wb = openpyxl.load_workbook(input_path)
    ws = wb["Hoja1"]

    rows = list(ws.iter_rows())

    # Find header row (row index 1, 0-indexed)
    header_row_idx = 1
    # Col indices (0-based): Cliente=0, Razon=1, Dom=2, Loc=3, Cluster=4, Contacto=5, Web=6
    COL_CONTACTO = 6  # column F (1-indexed) → col index 5 + 1
    COL_WEB = 7       # column G (1-indexed) → col index 6 + 1

    # Ensure headers exist
    ws.cell(row=2, column=6, value="Contacto")
    ws.cell(row=2, column=7, value="Web")

    data_rows = [(i, row) for i, row in enumerate(ws.iter_rows(min_row=3), start=3)
                 if row[0].value is not None and str(row[0].value).replace(".0", "").isdigit()]

    total = len(data_rows)
    found_count = 0
    errors = 0

    print(f"Total tiendas: {total}")
    print(f"Empezando desde fila: {start_row + 1}")
    print("-" * 60)

    for idx, (row_num, row) in enumerate(data_rows):
        if idx < start_row:
            continue

        razon = str(row[1].value or "").strip()
        domicilio = str(row[2].value or "").strip()
        localidad = str(row[3].value or "").strip()

        # Skip if already filled
        if row[5].value and row[6].value:
            print(f"[{idx+1}/{total}] Ya tiene datos, salteando: {razon}")
            continue

        try:
            result = find_place(razon, domicilio, localidad, api_key)

            phone = result["phone"] or ""
            website = result["website"] or ""

            ws.cell(row=row_num, column=6, value=phone)
            ws.cell(row=row_num, column=7, value=website)

            status = "✓" if result["found"] else "✗"
            if result["found"]:
                found_count += 1

            print(f"[{idx+1}/{total}] {status} {razon} | {localidad}")
            if phone or website:
                print(f"         Tel: {phone} | Web: {website}")

        except Exception as e:
            errors += 1
            print(f"[{idx+1}/{total}] ERROR {razon}: {e}")

        # Save every 50 rows to avoid losing progress
        if (idx + 1) % 50 == 0:
            wb.save(output_path)
            print(f"  → Guardado parcial en {output_path} (fila {idx+1}/{total})")

        time.sleep(delay)

    wb.save(output_path)
    print("\n" + "=" * 60)
    print(f"Completado. Encontrados: {found_count}/{total} | Errores: {errors}")
    print(f"Resultado guardado en: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Buscar contactos de puntos de venta en Argentina")
    parser.add_argument("--api-key", required=True, help="Google Places API key")
    parser.add_argument("--input", required=True, help="Archivo Excel de entrada (.xlsx)")
    parser.add_argument("--output", required=True, help="Archivo Excel de salida (.xlsx)")
    parser.add_argument("--delay", type=float, default=0.3, help="Segundos entre requests (default: 0.3)")
    parser.add_argument("--start-row", type=int, default=0, help="Fila desde la que empezar (para reanudar)")

    args = parser.parse_args()

    if not Path(args.input).exists():
        print(f"Error: No se encontró el archivo {args.input}")
        sys.exit(1)

    import shutil
    if args.input != args.output:
        shutil.copy2(args.input, args.output)

    process_excel(args.input, args.output, args.api_key, args.delay, args.start_row)


if __name__ == "__main__":
    main()
